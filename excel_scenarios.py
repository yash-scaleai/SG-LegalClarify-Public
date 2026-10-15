from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_EXCEL_PATH = Path("SG-LegalClarify_public.xlsx")
DEFAULT_SHEET_NAME = "main"


def _to_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower())
    return re.sub(r"_+", "_", slug).strip("_")


def _normalize_importance(raw: str) -> str:
    text = raw.strip().lower()
    if text in {"mandatory", "essential"}:
        return "Mandatory"
    return "Optional"


def _normalize_question_type(raw: str) -> str:
    text = raw.strip().lower()
    if text == "fully specified":
        return "Fully specified"
    if text == "underspecified":
        return "Underspecified"
    return raw.strip() or "Unknown"


def _resolve_question_text(
    row: tuple[Any, ...],
    header_idx: dict[str, int],
    question_type_raw: str,
) -> str:
    """
    Resolve scenario question text across legacy and split question schemas.

    Preferred mapping:
    - QUESTION_TYPE == underspecified   -> UNDERSPECIFIED QUESTION
    - QUESTION_TYPE == fully specified  -> FULLY SPECIFIED QUESTION

    Falls back to legacy QUESTION and then whichever split column is populated.
    """
    question_type = question_type_raw.strip().lower()
    underspecified_question = _to_str(row[header_idx["UNDERSPECIFIED QUESTION"]]) if "UNDERSPECIFIED QUESTION" in header_idx else ""
    fully_specified_question = _to_str(row[header_idx["FULLY SPECIFIED QUESTION"]]) if "FULLY SPECIFIED QUESTION" in header_idx else ""
    legacy_question = _to_str(row[header_idx["QUESTION"]]) if "QUESTION" in header_idx else ""

    if question_type == "underspecified":
        return underspecified_question or legacy_question or fully_specified_question
    if question_type == "fully specified":
        return fully_specified_question or legacy_question or underspecified_question

    # Unknown/missing type: prefer legacy, then whichever split column has text.
    return legacy_question or underspecified_question or fully_specified_question


def _parse_clarifying_variables(raw: str) -> dict[str, dict[str, str]]:
    """
    Parse CLARIFYING_QUESTION text formatted as:
    Q: ... A: ... | Q: ... A: ...
    """
    text = raw.strip()
    if not text:
        return {}

    pattern = re.compile(r"Q:\s*(.*?)\s*A:\s*(.*?)(?=\s*\|\s*Q:|$)", re.IGNORECASE | re.DOTALL)
    matches = pattern.findall(text)
    variables: dict[str, dict[str, str]] = {}
    for i, (question, answer) in enumerate(matches, start=1):
        q = question.strip()
        a = answer.strip()
        if not q and not a:
            continue
        variables[f"var_{i-1}"] = {
            "question": q,
            "answer": a,
        }
    return variables


def _derive_title(question: str, fallback: str) -> str:
    cleaned = re.sub(r"\s+", " ", question).strip()
    if not cleaned:
        return fallback
    short = cleaned[:72].rstrip(" .,!?:;")
    return short


def _discover_criterion_indices(header_idx: dict[str, int]) -> list[int]:
    indices: set[int] = set()
    for header_name in header_idx:
        match = re.match(r"^CRITERION_(\d+)(?:_(?:DIMENSION|ESSENTIAL))?$", header_name)
        if match:
            indices.add(int(match.group(1)))
    return sorted(indices)


def _build_converted_fully_specified_scenario(
    base_scenario: dict[str, Any],
    fully_specified_question: str,
    scenarios: dict[str, dict],
) -> tuple[str, dict] | None:
    """
    Build a synthetic fully-specified scenario for underspecified rows
    that also define FULLY SPECIFIED QUESTION.
    """
    question = fully_specified_question.strip()
    if not question:
        return None

    base_task_id = str(base_scenario.get("task_id", "")).strip()
    if not base_task_id:
        return None

    candidate_key = f"{base_task_id}__fully_specified"
    if candidate_key in scenarios:
        suffix = 2
        while f"{candidate_key}_{suffix}" in scenarios:
            suffix += 1
        candidate_key = f"{candidate_key}_{suffix}"

    converted = dict(base_scenario)
    converted.update(
        {
            "task_id": candidate_key,
            "title": _derive_title(question, base_scenario.get("subtopic", "scenario")),
            "output_slug": _slugify(candidate_key),
            "question_type": "Fully specified",
            "user_question": question,
            "variables": {},
            "is_converted_fully_specified": True,
            "converted_from_task_id": base_task_id,
        }
    )
    return candidate_key, converted


def load_scenarios_from_excel(
    excel_path: Path | str = DEFAULT_EXCEL_PATH,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> dict[str, dict]:
    workbook_path = Path(excel_path)
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[sheet_name]

    headers = [_to_str(c.value) for c in ws[1]]
    header_idx = {h: i for i, h in enumerate(headers) if h}
    # Accept both space and underscore variants of any column name (e.g.
    # 'UNDERSPECIFIED QUESTION' vs 'UNDERSPECIFIED_QUESTION') so this loader is
    # robust across sheet exports that switch conventions. Only add an alias when
    # it doesn't collide with a distinct existing header.
    for name, idx in list(header_idx.items()):
        for alias in {name.replace("_", " "), name.replace(" ", "_")}:
            if alias != name and alias not in header_idx:
                header_idx[alias] = idx
    criterion_indices = _discover_criterion_indices(header_idx)

    scenarios: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        task_id = _to_str(row[header_idx.get("TASK_ID", -1)]) or f"task_{len(scenarios)+1}"
        question_type_raw = _to_str(row[header_idx.get("QUESTION_TYPE", -1)])
        question = _resolve_question_text(row, header_idx, question_type_raw)
        if not question:
            continue

        underspecified_question = _to_str(row[header_idx.get("UNDERSPECIFIED QUESTION", -1)])
        fully_specified_question = _to_str(row[header_idx.get("FULLY SPECIFIED QUESTION", -1)])
        question_type = question_type_raw.lower()
        category = _to_str(row[header_idx.get("CATEGORY", -1)]) or "Unknown"
        topic = _to_str(row[header_idx.get("TOPIC", -1)]) or "Unknown"
        subtopic = _to_str(row[header_idx.get("SUBTOPIC", -1)]) or "scenario"

        clarifying_blob = _to_str(row[header_idx.get("CLARIFYING_QUESTION", -1)])
        variables = _parse_clarifying_variables(clarifying_blob) if question_type == "underspecified" else {}

        rubric = []
        for i in criterion_indices:
            item_text = _to_str(row[header_idx.get(f"CRITERION_{i}", -1)])
            if not item_text:
                continue
            dimension = _to_str(row[header_idx.get(f"CRITERION_{i}_DIMENSION", -1)]) or "Issue"
            essential = _to_str(row[header_idx.get(f"CRITERION_{i}_ESSENTIAL", -1)])
            rubric.append(
                {
                    "id": i,
                    "item": item_text,
                    "dimension": dimension,
                    "importance": _normalize_importance(essential),
                }
            )

        scenario_key = task_id
        title = _derive_title(question, subtopic)
        scenarios[scenario_key] = {
            "task_id": task_id,
            "title": title,
            "output_slug": _slugify(task_id),
            "category": category,
            "topic": topic,
            "subtopic": subtopic,
            "question_type": _normalize_question_type(question_type_raw),
            "user_question": question,
            "variables": variables,
            "rubric": rubric,
            "underspecified_question": underspecified_question,
            "fully_specified_question": fully_specified_question,
            "is_converted_fully_specified": False,
        }

        # If an underspecified row also has a fully specified question, emit
        # a second synthetic scenario that is treated as a normal fully-specified run.
        if question_type == "underspecified":
            converted_result = _build_converted_fully_specified_scenario(
                scenarios[scenario_key],
                fully_specified_question=fully_specified_question,
                scenarios=scenarios,
            )
            if converted_result is not None:
                converted_key, converted_scenario = converted_result
                scenarios[converted_key] = converted_scenario

    return scenarios

